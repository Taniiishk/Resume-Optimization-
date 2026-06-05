"""Resume parser: PDF / DOCX / TXT -> structured Resume model.

Strategy:
1. Extract plain text from the upload.
2. Send the text to Gemini with a strict JSON schema prompt to get
   structured fields (contact, summary, skills, experience, etc.).
3. Validate and return a Resume model.
"""

from __future__ import annotations

import io
import logging
import re

import pdfplumber
from docx import Document

from app.core.llm import generate_json
from app.models.schemas import Resume

logger = logging.getLogger(__name__)


# ---------- Text extraction ----------

def _clean_text(s: str) -> str:
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def extract_text_from_pdf(data: bytes) -> str:
    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            txt = page.extract_text() or ""
            if txt:
                parts.append(txt)
    return _clean_text("\n\n".join(parts))


def extract_text_from_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        logger.warning("Failed to open DOCX: %s — falling back to text extraction", e)
        return extract_text_from_txt(data)
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    parts.append(cell.text.strip())
    return _clean_text("\n".join(parts))


def extract_text_from_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="ignore")
    text = _clean_text(text)
    printable = sum(1 for c in text if c.isprintable() or c in "\n\r\t")
    if len(text) > 0 and printable / max(len(text), 1) < 0.8:
        raise ValueError("File does not appear to be a text document (binary content detected)")
    return text


def extract_text(filename: str, content_type: str, data: bytes) -> str:
    # Reject files that are actually images (magic bytes)
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:2] in (b"\xff\xd8", b"BM"):
        raise ValueError("File appears to be an image, not a text document")
    name = (filename or "").lower()
    if name.endswith(".pdf") or "pdf" in (content_type or "").lower():
        return extract_text_from_pdf(data)
    if name.endswith(".docx") or "word" in (content_type or "").lower():
        return extract_text_from_docx(data)
    if name.endswith((".txt", ".text", ".md", ".csv", ".json", ".xml", ".rtf")):
        return extract_text_from_txt(data)
    # Unknown extension — try PDF first, then TXT
    try:
        return extract_text_from_pdf(data)
    except Exception:
        pass
    return extract_text_from_txt(data)


# ---------- LLM-based structuring ----------

_PARSE_PROMPT = """You are a precise resume-parsing engine. Convert the following resume text into a strict JSON object.

Rules:
- Output ONLY valid JSON. No commentary, no markdown fences.
- Use empty strings or empty arrays for fields you cannot find.
- Normalize the contact info (trim whitespace).
- For `skills`, return a flat array of individual skills, not categories.
- For experience `bullets`, keep each bullet as a separate string, trimmed.
- Preserve dates exactly as written in the resume.

Required JSON schema:
{{
  "contact": {{
    "name": "",
    "email": "",
    "phone": "",
    "location": "",
    "linkedin": "",
    "github": "",
    "portfolio": ""
  }},
  "summary": "",
  "skills": [],
  "experience": [
    {{
      "title": "",
      "company": "",
      "location": "",
      "start_date": "",
      "end_date": "",
      "bullets": []
    }}
  ],
  "education": [
    {{
      "degree": "",
      "school": "",
      "location": "",
      "start_date": "",
      "end_date": "",
      "gpa": "",
      "details": []
    }}
  ],
  "projects": [
    {{
      "name": "",
      "link": "",
      "tech": [],
      "bullets": []
    }}
  ],
  "certifications": [
    {{
      "name": "",
      "issuer": "",
      "date": ""
    }}
  ]
}}

Resume text:
\"\"\"
{resume_text}
\"\"\"
"""


def parse_resume_text(resume_text: str) -> Resume:
    if not resume_text or not resume_text.strip():
        raise ValueError("Resume text is empty")

    prompt = _PARSE_PROMPT.format(resume_text=resume_text[:12000])  # safety cap
    raw = generate_json(prompt)

    # Normalize list-shaped fields defensively
    if isinstance(raw.get("skills"), str):
        raw["skills"] = [s.strip() for s in re.split(r"[,;\n|/]", raw["skills"]) if s.strip()]

    for job in raw.get("experience", []) or []:
        if isinstance(job.get("bullets"), str):
            job["bullets"] = [b.strip(" \t-•") for b in job["bullets"].split("\n") if b.strip()]

    for proj in raw.get("projects", []) or []:
        if isinstance(proj.get("tech"), str):
            proj["tech"] = [t.strip() for t in re.split(r"[,;|]", proj["tech"]) if t.strip()]
        if isinstance(proj.get("bullets"), str):
            proj["bullets"] = [b.strip(" \t-•") for b in proj["bullets"].split("\n") if b.strip()]

    for edu in raw.get("education", []) or []:
        if isinstance(edu.get("details"), str):
            edu["details"] = [d.strip() for d in edu["details"].split("\n") if d.strip()]

    raw["raw_text"] = resume_text
    return Resume.model_validate(raw)


def parse_resume_upload(filename: str, content_type: str, data: bytes) -> Resume:
    text = extract_text(filename, content_type, data)
    logger.info("Extracted %d chars from %s", len(text), filename)
    return parse_resume_text(text)
