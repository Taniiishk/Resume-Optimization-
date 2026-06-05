from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.models.schemas import Resume

BLUE = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x1E, 0x29, 0x3B)
GRAY = RGBColor(0x47, 0x55, 0x69)


def _add_section_line(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single",
        qn("w:sz"): "6",
        qn("w:space"): "1",
        qn("w:color"): "2563EB",
    })
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_run(paragraph, text: str, bold=False, italic=False, color=None, size=11):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return run


def render_resume_docx(resume: Resume) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

    # Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    _add_run(p, resume.contact.name or "Resume", bold=True, size=22, color=DARK)

    # Contact
    parts: list[str] = []
    c = resume.contact
    if c.location:
        parts.append(c.location)
    if c.phone:
        parts.append(c.phone)
    if c.email:
        parts.append(c.email)
    if c.linkedin:
        parts.append(c.linkedin)
    if c.github:
        parts.append(c.github)
    if c.portfolio:
        parts.append(c.portfolio)
    if parts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        _add_run(p, "  |  ".join(parts), size=9, color=GRAY)

    # Helper
    def add_section(title: str):
        p = doc.add_paragraph()
        _add_run(p, title, bold=True, size=12, color=DARK)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        _add_section_line(doc)

    def add_heading_line(left: str, right: str = ""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        _add_run(p, left, bold=True, size=11, color=DARK)
        if right:
            _add_run(p, "   ", size=11)
            _add_run(p, right, italic=True, size=10, color=GRAY)

    def add_italic_line(text: str):
        if not text:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        _add_run(p, text, italic=True, size=10, color=GRAY)

    def add_bullet(text: str):
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        _add_run(p, text, size=10.5, color=DARK)

    # Summary
    if resume.summary and resume.summary.strip():
        add_section("Professional Summary")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _add_run(p, resume.summary.strip(), size=11, color=DARK)

    # Skills
    if resume.skills:
        add_section("Skills")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _add_run(p, "  •  ".join(s for s in resume.skills if s), size=10.5, color=DARK)

    # Experience
    if resume.experience:
        add_section("Experience")
        for job in resume.experience:
            if not job.title and not job.company:
                continue
            left = " — ".join(filter(None, [job.title, job.company]))
            right = ""
            if job.start_date or job.end_date:
                parts = [p for p in [job.start_date, job.end_date] if p]
                right = " – ".join(parts)
            add_heading_line(left, right)
            add_italic_line(job.location)
            if job.bullets:
                for b in job.bullets:
                    if b.strip():
                        add_bullet(b.strip())

    # Projects
    if resume.projects:
        add_section("Projects")
        for pj in resume.projects:
            if not pj.name:
                continue
            left = pj.name
            right = pj.link if pj.link else ""
            add_heading_line(left, right)
            if pj.tech:
                add_italic_line(f"Technologies: {', '.join(pj.tech)}")
            if pj.bullets:
                for b in pj.bullets:
                    if b.strip():
                        add_bullet(b.strip())

    # Education
    if resume.education:
        add_section("Education")
        for edu in resume.education:
            if not edu.degree and not edu.school:
                continue
            left = " — ".join(filter(None, [edu.degree, edu.school]))
            right = ""
            if edu.start_date or edu.end_date:
                parts = [p for p in [edu.start_date, edu.end_date] if p]
                right = " – ".join(parts)
            add_heading_line(left, right)
            details = []
            if edu.location:
                details.append(edu.location)
            if edu.gpa:
                details.append(f"GPA: {edu.gpa}")
            add_italic_line(" | ".join(details))
            if edu.details:
                for d in edu.details:
                    if d.strip():
                        add_bullet(d.strip())

    # Certifications
    if resume.certifications:
        add_section("Certifications")
        for cert in resume.certifications:
            if cert.name:
                parts = [cert.name]
                if cert.issuer:
                    parts.append(cert.issuer)
                if cert.date:
                    parts.append(f"({cert.date})")
                p = doc.add_paragraph(style="List Bullet")
                p.clear()
                _add_run(p, " — ".join(parts), size=10.5, color=DARK)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def safe_filename(name: str) -> str:
    name = (name or "resume").strip()
    name = re.sub(r"[^A-Za-z0-9_-]+", "_", name)
    return name
